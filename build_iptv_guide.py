from __future__ import annotations

import math
import os
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "iptv_guide_assets"
OUT = ROOT / "IPTV_End_to_End_Field_Guide.docx"
ASSET_DIR.mkdir(exist_ok=True)


# compact_reference_guide preset, plus named editorial-cover and technical-diagram overrides.
PAGE_W_IN = 8.5
PAGE_H_IN = 11.0
MARGIN_IN = 1.0
CONTENT_W_IN = 6.5
CONTENT_W_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEAL = "148A8A"
GOLD = "B4872D"
INK = "1D2733"
MUTED = "5D6A77"
LIGHT_BLUE = "E8EEF5"
LIGHT_TEAL = "E8F4F3"
LIGHT_GOLD = "F7F0DE"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D7DEE6"
WHITE = "FFFFFF"
RED = "9B1C1C"
GREEN = "2F6B4F"


SOURCES = {
    "S1": ("Wikipedia: Internet Protocol television", "https://en.wikipedia.org/wiki/Internet_Protocol_television", "Background terminology and managed-IPTV architecture."),
    "S2": ("RFC 8216: HTTP Live Streaming", "https://www.rfc-editor.org/rfc/rfc8216.html", "HLS playlists, variants, media segments, MPEG-TS, and fragmented MP4."),
    "S3": ("RFC 3550: RTP", "https://www.rfc-editor.org/rfc/rfc3550.html", "RTP sequence numbers, timestamps, RTCP, and interarrival jitter."),
    "S4": ("RFC 7826: RTSP 2.0", "https://www.rfc-editor.org/rfc/rfc7826.html", "Session description, SETUP, PLAY, PAUSE, and media transport control."),
    "S5": ("RFC 9776: IGMPv3", "https://www.rfc-editor.org/rfc/rfc9776.html", "Current IGMPv3 group and source membership behavior; obsoletes RFC 3376."),
    "S6": ("RFC 4541: IGMP and MLD Snooping Switches", "https://www.rfc-editor.org/rfc/rfc4541.html", "Layer-2 multicast forwarding considerations."),
    "S7": ("DASH-IF: Adaptive Bitrate Streaming", "https://dashif.org/dash.js/pages/usage/abr/", "Throughput-, buffer-, and device-aware representation selection."),
    "S8": ("XMLTV DTD", "https://github.com/XMLTV/xmltv/blob/master/xmltv.dtd", "Normative XMLTV structure, channel identity, time format, and programme metadata."),
    "S9": ("XMLTV File Format", "https://wiki.xmltv.org/index.php/XMLTVFormat", "Readable overview and examples of XMLTV channel and programme records."),
    "S10": ("Kodi IPTV Simple client", "https://github.com/kodi-pvr/pvr.iptvsimple", "Widely implemented extended-M3U fields, groups, EPG mapping, and catch-up conventions."),
    "S11": ("HDHomeRun technical documentation", "https://info.hdhomerun.com/info/tech", "Tuning, UDP targets, HTTP streaming, and lineup endpoints."),
    "S12": ("HDHomeRun Device Discover API", "https://info.hdhomerun.com/info/discovery_api", "Discovery metadata, tuner count, base URL, and lineup URL."),
    "S13": ("Jellyfin Live TV Setup Guide", "https://jellyfin.org/docs/general/server/live-tv/setup-guide/", "HDHomeRun/M3U tuner configuration, XMLTV, mapping, and stream limits."),
    "S14": ("Jellyfin Transcoding", "https://jellyfin.org/docs/general/post-install/transcoding/", "Direct Play, Remux, Direct Stream, Transcode, and client capability selection."),
    "S15": ("Jellyfin Codec Support", "https://jellyfin.org/docs/general/clients/codec-support/", "Client/container compatibility and common remux/transcode triggers."),
    "S16": ("Jellyfin Hardware Acceleration", "https://jellyfin.org/docs/general/post-install/transcoding/hardware-acceleration/", "jellyfin-ffmpeg and supported hardware acceleration paths."),
    "S17": ("FFmpeg documentation", "https://ffmpeg.org/ffmpeg.html", "Streamcopy, transcoding, filtering, stream selection, and mapping."),
    "S18": ("FFmpeg protocols", "https://ffmpeg.org/ffmpeg-protocols.html", "HTTP reconnects, timeouts, UDP buffers, RTP/RTSP, and SRT options."),
    "S19": ("FFmpeg formats", "https://ffmpeg.org/ffmpeg-formats.html", "Demuxer/muxer options, probing, MPEG-TS, HLS, and segmentation."),
    "S20": ("ffprobe documentation", "https://ffmpeg.org/ffprobe.html", "Machine-readable inspection of formats, programs, streams, packets, and frames."),
    "S21": ("VLC core input options", "https://github.com/videolan/vlc/blob/master/src/libvlc-module.c", "Network/live/file caching, clock jitter, synchronization, and late-frame controls."),
    "S22": ("VLC clock architecture", "https://github.com/videolan/vlc/blob/master/doc/clock.md", "Synchronization of elementary streams and selectable clock masters."),
    "S23": ("VLC: Stream over HTTP", "https://docs.videolan.me/vlc-user/desktop/3.0/en/advanced/streaming/stream_over_http.html", "HTTP pull streaming and MPEG-TS stream-output examples."),
    "S24": ("VideoLAN: VLC Features", "https://www.videolan.org/vlc/features.html", "Protocol, container, codec, subtitle, and hardware-decoding coverage."),
    "S25": ("Threadfin", "https://github.com/Threadfin/Threadfin", "M3U/XMLTV merging, filtering, mapping, backup streams, categories, and tuner limits."),
    "S26": ("Tvheadend Electronic Program Guide", "https://docs.tvheadend.org/documentation/configuration/electronic-program-guide", "OTA/XMLTV guide ingestion, autorecording, DVR profiles, and padding."),
    "S27": ("Tvheadend DVR API", "https://docs.tvheadend.org/documentation/development/json-api/api-description/dvr", "Timers, series rules, real start/stop times, and recorder padding."),
    "S28": ("ErsatzTV Collections", "https://ersatztv.org/docs/collections/", "Dynamic smart collections used as schedulable content sources."),
    "S29": ("ErsatzTV Playout Instructions", "https://ersatztv.org/docs/scheduling/sequential/playout/", "Sequential schedule construction, filler, loops, and EPG grouping."),
    "S30": ("dizqueTV", "https://github.com/vexorian/dizquetv", "Virtual linear channels, XMLTV output, HDHomeRun emulation, and optional transcoding."),
    "S31": ("Channels DVR: Custom Channels", "https://getchannels.com/docs/channels-dvr-server/how-to/custom-channels/", "M3U/XMLTV custom channels, source identifiers, and supported stream types."),
    "S32": ("Wikipedia: MPEG transport stream", "https://en.wikipedia.org/wiki/MPEG_transport_stream", "MPEG-TS packets, PIDs, PAT/PMT, PCR, PTS, and multiplexing context."),
    "S33": ("ITU-T H.264", "https://www.itu.int/rec/T-REC-H.264/en", "Current H.264/AVC recommendation and application scope."),
    "S34": ("Alliance for Open Media: AV1", "https://aomedia.org/specifications/av1/", "AV1 bitstream specification and ecosystem overview."),
    "S35": ("Matroska Data Layout", "https://www.matroska.org/technical/diagram.html", "Matroska segments, tracks, clusters, and timing structure."),
    "S36": ("W3C Encrypted Media Extensions", "https://www.w3.org/TR/encrypted-media-2/", "Browser APIs for key-system/CDM interaction and protected playback."),
    "S37": ("Apple FairPlay Streaming", "https://developer.apple.com/streaming/fps/", "Protected HLS delivery and license/key exchange on Apple platforms."),
    "S38": ("Haivision SRT", "https://github.com/Haivision/srt", "Low-latency reliable UDP transport, ARQ, encryption, and loss recovery."),
    "S39": ("ATSC A/65: Program and System Information Protocol", "https://www.atsc.org/wp-content/uploads/2021/04/A65_2013.pdf", "ATSC virtual-channel and EPG tables."),
    "S40": ("ETSI EN 300 468: DVB Service Information", "https://portal.etsi.org/webapp/ewp/copy_file.asp?wki_id=72198", "DVB service and event information, including schedule and present/following EIT."),
    "S41": ("iptv-org project", "https://github.com/iptv-org/iptv", "Example of curated public stream lineups, stable IDs, validation, and takedown workflow."),
    "S42": ("Wikipedia: Electronic program guide", "https://en.wikipedia.org/wiki/Electronic_program_guide", "Historical and cross-platform EPG context."),
}


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(kwargs.get(margin, CELL_MARGINS_DXA[margin])))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6", inside=True) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    edges = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for edge in edges:
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is not None:
        return
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_table_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_W_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_W_DXA}: {widths_dxa}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_W_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def add_hyperlink(paragraph, text: str, url: str, color=BLUE, underline=True):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def set_image_alt(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("name", title)
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_numbering(doc: Document, fmt: str, text: str, left: int, hanging: int) -> int:
    numbering = doc.part.numbering_part.element
    abs_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_id = max(abs_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abs_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_num = OxmlElement("w:abstractNumId")
    abs_num.set(qn("w:val"), str(abs_id))
    num.append(abs_num)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int, level: int = 0) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def configure_styles(doc: Document) -> tuple[int, int]:
    sec = doc.sections[0]
    sec.page_width = Inches(PAGE_W_IN)
    sec.page_height = Inches(PAGE_H_IN)
    sec.top_margin = Inches(MARGIN_IN)
    sec.bottom_margin = Inches(MARGIN_IN)
    sec.left_margin = Inches(MARGIN_IN)
    sec.right_margin = Inches(MARGIN_IN)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    sec.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(9)
    caption.paragraph_format.keep_with_next = True

    for name, base, size, color, before, after, line in [
        ("Lead", "Normal", 12.5, NAVY, 0, 10, 1.25),
        ("Small Text", "Normal", 9, MUTED, 0, 4, 1.15),
        ("Source Text", "Normal", 9, MUTED, 0, 5, 1.15),
        ("Table Text", "Normal", 9.3, INK, 0, 0, 1.1),
        ("Table Header", "Normal", 9.3, NAVY, 0, 0, 1.1),
        ("Callout Label", "Normal", 9, DARK_BLUE, 0, 2, 1.0),
        ("Code Block", "Normal", 8.5, INK, 0, 0, 1.0),
    ]:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles[base]
        style.font.name = "Consolas" if name == "Code Block" else "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = name in ("Table Header", "Callout Label")
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = line

    code_style = doc.styles["Code Block"]
    code_style.paragraph_format.left_indent = Inches(0.16)
    code_style.paragraph_format.right_indent = Inches(0.08)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(6)

    bullet_num_id = add_numbering(doc, "bullet", "•", 540, 270)
    decimal_num_id = add_numbering(doc, "decimal", "%1.", 540, 270)
    return bullet_num_id, decimal_num_id


def set_document_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "IPTV, End to End: A Technical Field Guide"
    props.subject = "How IPTV lineups, guide data, tuners, schedules, codecs, players, FFmpeg, VLC, and Jellyfin fit together"
    props.author = "OpenAI Codex"
    props.keywords = "IPTV, XMLTV, EPG, M3U, HDHomeRun, VLC, FFmpeg, Jellyfin, HLS, MPEG-TS"
    props.comments = "Research edition compiled from primary standards, official project documentation, and representative open-source implementations."


def setup_header_footer(doc: Document) -> None:
    sec = doc.sections[0]
    header = sec.header
    p = header.paragraphs[0]
    p.text = "IPTV, END TO END"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        set_run_font(run, size=8.5, color=MUTED, bold=True)
    footer = sec.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(CONTENT_W_IN))
    set_table_geometry(table, [6500, 2860], indent_dxa=0)
    set_repeat_table_header(table.rows[0])
    # The footer table is furniture; remove all visible borders.
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tbl_pr.append(borders)
    left = table.cell(0, 0).paragraphs[0]
    left.add_run("Research edition • August 2026")
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.add_run("Page ")
    add_field(right, "PAGE")
    right.add_run(" of ")
    add_field(right, "NUMPAGES")
    for cell in table.rows[0].cells:
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(0)
            for run in para.runs:
                set_run_font(run, size=8.5, color=MUTED)
    first_header = sec.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = sec.first_page_footer
    first_footer.paragraphs[0].text = ""


def add_title_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TECHNICAL FIELD GUIDE")
    set_run_font(r, size=10, color=GOLD, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("IPTV, End to End")
    set_run_font(r, size=32, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("How lineups, guide data, tuners, schedules, codecs, players, FFmpeg, VLC, and Jellyfin fit together")
    set_run_font(r, size=14, color=DARK_BLUE)

    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [3120, 3120, 3120], indent_dxa=TABLE_INDENT_DXA)
    set_repeat_table_header(table.rows[0])
    set_table_borders(table, color=WHITE, size="0", inside=False)
    labels = [
        ("CONTROL PLANE", "Providers • playlists • identity • EPG"),
        ("MEDIA PLANE", "Protocols • containers • codecs • timing"),
        ("OPERATIONS PLANE", "Tuners • recording • buffering • health"),
    ]
    for idx, (label, detail) in enumerate(labels):
        cell = table.cell(0, idx)
        set_cell_shading(cell, [LIGHT_BLUE, LIGHT_TEAL, LIGHT_GOLD][idx])
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(4)
        rr = p1.add_run(label)
        set_run_font(rr, size=8.5, color=[BLUE, TEAL, GOLD][idx], bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        rr = p2.add_run(detail)
        set_run_font(rr, size=9.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(44)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("A practical reference for builders, operators, and technically curious readers")
    set_run_font(r, size=11, color=MUTED, italic=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Research edition • 19 August 2026")
    set_run_font(r, size=10, color=NAVY, bold=True)
    doc.add_page_break()


def add_citations(paragraph, refs: str) -> None:
    run = paragraph.add_run(f" [{refs}]")
    set_run_font(run, size=8.5, color=BLUE, bold=True)


def add_body(doc: Document, text: str, refs: str | None = None, style: str | None = None, bold_lead: str | None = None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        first, rest = text[: len(bold_lead)], text[len(bold_lead):]
        r = p.add_run(first)
        r.bold = True
        p.add_run(rest)
    else:
        p.add_run(text)
    if refs:
        add_citations(p, refs)
    return p


def add_bullet(doc: Document, text: str, num_id: int, refs: str | None = None):
    p = doc.add_paragraph()
    apply_num(p, num_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.add_run(text)
    if refs:
        add_citations(p, refs)
    return p


def add_step(doc: Document, text: str, num_id: int, refs: str | None = None):
    p = doc.add_paragraph()
    apply_num(p, num_id)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    p.add_run(text)
    if refs:
        add_citations(p, refs)
    return p


def add_code(doc: Document, code: str, label: str | None = None) -> None:
    if label:
        p = doc.add_paragraph(style="Callout Label")
        p.paragraph_format.keep_with_next = True
        p.add_run(label.upper())
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    lines = code.strip("\n").splitlines()
    for idx, line in enumerate(lines):
        p.add_run(line if line else " ")
        if idx < len(lines) - 1:
            p.add_run().add_break()
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)


def add_callout(doc: Document, label: str, text: str, fill=LIGHT_BLUE, accent=BLUE, refs: str | None = None) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_W_DXA])
    set_table_borders(table, color=accent, size="10", inside=False)
    cell = table.cell(0, 0)
    set_repeat_table_header(table.rows[0])
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label.upper())
    set_run_font(r, size=9, color=accent, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    p.add_run(text)
    if refs:
        add_citations(p, refs)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int], caption: str | None = None, small=False) -> None:
    if caption:
        p = doc.add_paragraph(style="Caption")
        p.add_run(caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.style = doc.styles["Table Header"]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run(header)
    set_repeat_table_header(table.rows[0])
    keep_table_row_together(table.rows[0])
    for ridx, row in enumerate(rows):
        table_row = table.add_row()
        keep_table_row_together(table_row)
        cells = table_row.cells
        for idx, value in enumerate(row):
            if ridx % 2 == 1:
                set_cell_shading(cells[idx], "F9FAFB")
            p = cells[idx].paragraphs[0]
            p.style = doc.styles["Small Text" if small else "Table Text"]
            p.add_run(value)
    set_table_geometry(table, widths)
    set_table_borders(table)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)


def chapter(doc: Document, number: str, title: str, new_page=False) -> None:
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(f"{number}  {title}")
    r.bold = True


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    p.add_run("Guide map")
    add_body(
        doc,
        "Use this map to move from signal acquisition to metadata, playback, operations, and troubleshooting. Source markers such as [S12] connect claims to the linked bibliography.",
        style="Small Text",
    )
    add_table(
        doc,
        ["Foundations and metadata", "Playback, operations, and references"],
        [
            ["0  Executive mental model", "10  VLC playback and buffering"],
            ["1  What IPTV means", "11  FFmpeg probing, remuxing, and transcoding"],
            ["2  Sources and provider models", "12  Jellyfin live-TV stream handling"],
            ["3  Network tuners", "13  Dynamic groups and virtual channels"],
            ["4  Lineups, playlists, and M3U", "14  Reliability and observability"],
            ["5  XMLTV, EPG, and guide ingestion", "15  Security, rights, and privacy"],
            ["6  Channel IDs and identity", "16  Reference architecture"],
            ["7  Scheduling, recording, and time-shift", "17  Troubleshooting playbook"],
            ["8  Transport protocols", "A  Quick-reference tables"],
            ["9  Containers, codecs, and timestamps", "B  Glossary"],
            ["", "C  Sources and further reading"],
        ],
        [4680, 4680],
    )
    add_callout(
        doc,
        "Reading paths",
        "New to IPTV? Read 0–12 in order. Building a platform? Focus on 3–9 and 11–16. Debugging an incident? Start with 17, then jump to the subsystem chapter named by the symptom.",
    )


FONT_REG = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"


def font(size: int, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def wrapped(draw, text, box, font_obj, fill, align="center", spacing=8):
    x0, y0, x1, y1 = box
    max_width = x1 - x0 - 28
    words = text.split()
    lines = []
    line = ""
    for word in words:
        candidate = word if not line else f"{line} {word}"
        if draw.textbbox((0, 0), candidate, font=font_obj)[2] <= max_width:
            line = candidate
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    heights = [draw.textbbox((0, 0), line, font=font_obj)[3] for line in lines]
    total = sum(heights) + spacing * max(0, len(lines) - 1)
    y = y0 + (y1 - y0 - total) / 2
    for line, h in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=font_obj)
        w = bbox[2] - bbox[0]
        x = x0 + (x1 - x0 - w) / 2 if align == "center" else x0 + 16
        draw.text((x, y), line, font=font_obj, fill=fill)
        y += h + spacing


def arrow(draw, start, end, color=(46, 116, 181), width=7):
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 24
    for delta in (2.55, -2.55):
        p = (end[0] + length * math.cos(angle + delta), end[1] + length * math.sin(angle + delta))
        draw.line([end, p], fill=color, width=width)


def make_architecture_diagram(path: Path) -> None:
    img = Image.new("RGB", (2400, 1180), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "IPTV is two synchronized pipelines", font=font(58, True), fill=(23, 50, 77))
    d.text((70, 125), "The media pipeline moves bytes; the metadata pipeline tells the system what those bytes mean and when to record them.", font=font(29), fill=(93, 106, 119))
    media_boxes = [
        (80, 300, 420, 520, "SOURCE", "Antenna tuner\nManaged headend\nHTTP/HLS provider"),
        (500, 300, 840, 520, "NORMALIZE", "Discover / fetch\nFilter / map\nChoose route"),
        (920, 300, 1260, 520, "SERVER", "Open tuner\nRemux or transcode\nRecord / timeshift"),
        (1340, 300, 1680, 520, "DELIVER", "UDP/RTP\nHTTP-TS\nHLS / DASH"),
        (1760, 300, 2100, 520, "CLIENT", "Buffer\nDemux / decode\nRender A/V"),
    ]
    for i, (x0, y0, x1, y1, label, detail) in enumerate(media_boxes):
        fill = [(232, 238, 245), (232, 244, 243), (247, 240, 222), (232, 238, 245), (232, 244, 243)][i]
        d.rounded_rectangle((x0, y0, x1, y1), 24, fill=fill, outline=(46, 116, 181), width=4)
        d.text((x0 + 24, y0 + 22), label, font=font(29, True), fill=(46, 116, 181))
        wrapped(d, detail, (x0 + 12, y0 + 65, x1 - 12, y1 - 12), font(27), (29, 39, 51))
        if i < len(media_boxes) - 1:
            arrow(d, (x1 + 10, (y0 + y1) // 2), (media_boxes[i + 1][0] - 12, (y0 + y1) // 2))
    d.text((80, 570), "MEDIA PLANE", font=font(26, True), fill=(46, 116, 181))
    meta_boxes = [
        (280, 760, 680, 990, "LINEUP", "M3U / tuner lineup\nChannel keys and URLs"),
        (850, 760, 1250, 990, "IDENTITY JOIN", "tvg-id ↔ XMLTV channel@id\nInternal stable channel ID"),
        (1420, 760, 1820, 990, "GUIDE + SCHEDULER", "XMLTV programmes\nRules, conflicts, padding"),
    ]
    for i, (x0, y0, x1, y1, label, detail) in enumerate(meta_boxes):
        d.rounded_rectangle((x0, y0, x1, y1), 24, fill=(247, 240, 222), outline=(180, 135, 45), width=4)
        d.text((x0 + 24, y0 + 24), label, font=font(28, True), fill=(180, 135, 45))
        wrapped(d, detail, (x0 + 12, y0 + 70, x1 - 12, y1 - 12), font(27), (29, 39, 51))
        if i < len(meta_boxes) - 1:
            arrow(d, (x1 + 10, (y0 + y1) // 2), (meta_boxes[i + 1][0] - 12, (y0 + y1) // 2), color=(180, 135, 45))
    d.text((80, 1030), "METADATA + CONTROL PLANE", font=font(26, True), fill=(180, 135, 45))
    arrow(d, (1620, 750), (1180, 535), color=(20, 138, 138), width=7)
    img.save(path, quality=95)


def make_identity_diagram(path: Path) -> None:
    img = Image.new("RGB", (2300, 1220), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "One channel, many identifiers", font=font(58, True), fill=(23, 50, 77))
    d.text((70, 125), "Treat each identifier as a key in a different namespace. Persist explicit mappings between them.", font=font(29), fill=(93, 106, 119))
    rows = [
        ("PHYSICAL / RF", "frequency, modulation, tuner index", "How hardware acquires a multiplex"),
        ("TRANSPORT", "transport_stream_id, service_id/program_number, PID", "How MPEG-TS identifies programs and elementary streams"),
        ("LINEUP", "M3U tvg-id, tvg-chno, source URL", "How an IPTV client discovers and labels a service"),
        ("GUIDE", "XMLTV channel@id", "How programmes attach to the service"),
        ("APPLICATION", "Jellyfin/TV backend internal ID or UUID", "How settings, favorites, recordings, and permissions persist"),
        ("DISPLAY", "name, logo, channel number, group", "What the user sees; mutable and not a safe key"),
    ]
    y = 245
    for idx, (layer, key, purpose) in enumerate(rows):
        fill = (232, 238, 245) if idx % 2 == 0 else (248, 250, 252)
        d.rounded_rectangle((90, y, 2210, y + 130), 18, fill=fill, outline=(215, 222, 230), width=3)
        d.text((125, y + 26), layer, font=font(27, True), fill=(46, 116, 181))
        wrapped(d, key, (530, y + 10, 1260, y + 120), font(27), (29, 39, 51), align="left")
        wrapped(d, purpose, (1280, y + 10, 2160, y + 120), font(25), (93, 106, 119), align="left")
        if idx < len(rows) - 1:
            arrow(d, (1150, y + 130), (1150, y + 155), color=(20, 138, 138), width=5)
        y += 155
    d.rounded_rectangle((330, 1120, 1970, 1190), 16, fill=(247, 240, 222), outline=(180, 135, 45), width=3)
    wrapped(d, "Rule: URL, display name, and channel number are attributes. A stable canonical ID is identity.", (350, 1125, 1950, 1185), font(26, True), (122, 90, 0))
    img.save(path, quality=95)


def make_playback_diagram(path: Path) -> None:
    img = Image.new("RGB", (2300, 1240), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "The server's playback decision", font=font(58, True), fill=(23, 50, 77))
    d.text((70, 125), "Compatibility is a conjunction: protocol + container + every selected track + subtitles + bitrate constraints.", font=font(29), fill=(93, 106, 119))
    boxes = {
        "source": (780, 230, 1520, 360, "SOURCE STREAM\nprotocol • container • tracks • bitrate"),
        "all": (780, 450, 1520, 580, "Can the client accept everything exactly as-is?"),
        "direct": (100, 690, 570, 840, "DIRECT PLAY\nNo media transformation\nLowest server cost"),
        "container": (680, 690, 1150, 840, "REMUX\nChange container/segmentation\nCopy audio + video"),
        "audio": (1260, 690, 1730, 840, "DIRECT STREAM\nCopy video\nTranscode audio"),
        "video": (1840, 690, 2210, 840, "TRANSCODE\nDecode + filter + encode\nHighest server cost"),
    }
    for name, (x0, y0, x1, y1, label) in boxes.items():
        if name in ("source", "all"):
            fill, outline = (232, 238, 245), (46, 116, 181)
        elif name == "direct":
            fill, outline = (232, 244, 243), (20, 138, 138)
        elif name == "video":
            fill, outline = (250, 232, 232), (155, 28, 28)
        else:
            fill, outline = (247, 240, 222), (180, 135, 45)
        d.rounded_rectangle((x0, y0, x1, y1), 24, fill=fill, outline=outline, width=4)
        wrapped(d, label, (x0 + 10, y0 + 8, x1 - 10, y1 - 8), font(28, True if name not in ("source", "all") else False), (29, 39, 51))
    arrow(d, (1150, 360), (1150, 445))
    for key in ("direct", "container", "audio", "video"):
        x0, y0, x1, y1, _ = boxes[key]
        arrow(d, (1150, 580), ((x0 + x1) // 2, y0 - 8), color=(93, 106, 119), width=5)
    labels = [(330, 630, "YES"), (910, 630, "container only"), (1490, 630, "audio/subtitle"), (2040, 630, "video/bitrate")]
    for x, y, text in labels:
        d.text((x - 70, y), text, font=font(22, True), fill=(93, 106, 119))
    d.rounded_rectangle((260, 980, 2040, 1140), 20, fill=(242, 244, 247), outline=(215, 222, 230), width=3)
    wrapped(d, "Subtitles are a frequent surprise: text may remux; image subtitles may need burning, which forces video transcoding. HLS segment boundaries also need usable keyframes.", (290, 995, 2010, 1125), font(27), (29, 39, 51))
    img.save(path, quality=95)


def make_buffer_diagram(path: Path) -> None:
    img = Image.new("RGB", (2300, 1020), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Buffering is a chain, not one bucket", font=font(58, True), fill=(23, 50, 77))
    d.text((70, 125), "A stall occurs when the downstream stage consumes faster than upstream can replenish usable, timestamped media.", font=font(29), fill=(93, 106, 119))
    boxes = [
        (90, 320, 430, 560, "NETWORK", "socket receive buffer\nTCP or UDP arrival"),
        (520, 320, 860, 560, "ACCESS", "HTTP/HLS fetch\nreconnect / retry"),
        (950, 320, 1290, 560, "DEMUX", "probe + packet queue\nPAT/PMT + timestamps"),
        (1380, 320, 1720, 560, "DECODE", "compressed packets\nframes + audio samples"),
        (1810, 320, 2150, 560, "OUTPUT", "A/V clocks\nrender + late-frame policy"),
    ]
    for i, (x0, y0, x1, y1, title, detail) in enumerate(boxes):
        d.rounded_rectangle((x0, y0, x1, y1), 22, fill=(232, 238, 245) if i % 2 == 0 else (232, 244, 243), outline=(46, 116, 181), width=4)
        d.text((x0 + 22, y0 + 22), title, font=font(28, True), fill=(46, 116, 181))
        wrapped(d, detail, (x0 + 10, y0 + 65, x1 - 10, y1 - 15), font(27), (29, 39, 51))
        if i < len(boxes) - 1:
            arrow(d, (x1 + 8, 440), (boxes[i + 1][0] - 8, 440))
    d.rounded_rectangle((250, 710, 2050, 900), 24, fill=(247, 240, 222), outline=(180, 135, 45), width=4)
    wrapped(d, "Larger buffers absorb more jitter and brief outages, but add startup delay, channel-change latency, memory use, and live-edge distance. They cannot repair a bad average bitrate, invalid timestamps, missing keyframes, or unsupported codecs.", (280, 725, 2020, 885), font(28), (29, 39, 51))
    img.save(path, quality=95)


def add_figure(doc: Document, path: Path, caption: str, alt: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(6.35))
    set_image_alt(shape, caption, alt)
    cp = doc.add_paragraph(style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.add_run(caption)


def build_document() -> Document:
    doc = Document()
    bullet_id, decimal_id = configure_styles(doc)
    set_document_properties(doc)
    setup_header_footer(doc)
    make_architecture_diagram(ASSET_DIR / "architecture.png")
    make_identity_diagram(ASSET_DIR / "identity.png")
    make_playback_diagram(ASSET_DIR / "playback_decision.png")
    make_buffer_diagram(ASSET_DIR / "buffer_chain.png")

    add_title_cover(doc)
    add_toc(doc)

    chapter(doc, "0", "Executive mental model", new_page=True)
    add_body(doc, "IPTV is not one protocol or one product. It is a coordinated system that acquires a live service, names it, attaches schedule metadata, transports compressed media, allocates finite resources, and adapts playback to the receiving device.", style="Lead")
    add_figure(doc, ASSET_DIR / "architecture.png", "Figure 1. The media plane and metadata/control plane meet at the TV server.", "Five media stages from source to client above a three-stage lineup, identity, guide and scheduler pipeline.")
    add_callout(doc, "The single most useful rule", "Keep four concepts separate: identity (which service?), location (which URL or tuner path now?), presentation (what name/number/logo?), and schedule (what airs when?). Most brittle IPTV systems collapse them into one mutable playlist line.", fill=LIGHT_TEAL, accent=TEAL)
    add_body(doc, "A channel change is a distributed transaction. The system selects a route, reserves a tuner or provider connection, opens the stream, recognizes its container and tracks, waits for decodable timing and a keyframe, chooses direct play/remux/transcode, fills enough buffer, then renders. The guide can be correct while playback fails, and playback can work while the guide is empty.")
    add_table(doc, ["Plane", "Core question", "Typical artifacts"], [
        ["Acquisition", "Where do the bytes originate?", "RF tuner, provider URL, RTSP session, virtual playout"],
        ["Catalog / identity", "Which logical channel is this?", "Canonical ID, M3U tvg-id, service ID, internal UUID"],
        ["Guide / schedule", "What is on, and what should be recorded?", "XMLTV, DVB EIT, ATSC PSIP, rules, padding"],
        ["Transport", "How do packets reach the consumer?", "UDP/RTP, HTTP-TS, HLS, DASH, SRT"],
        ["Media", "How are pictures, sound, subtitles, and clocks represented?", "MPEG-TS/fMP4/MKV; H.264/HEVC; AAC/AC-3; PTS/PCR"],
        ["Playback", "Can this client consume the result?", "Direct play, remux, audio transcode, video transcode"],
        ["Operations", "Will it stay reliable?", "Health probes, concurrency, logs, metrics, retries, failover"],
    ], [1700, 3000, 4660], caption="Table 1. A seven-plane model for reasoning about IPTV systems.")

    chapter(doc, "1", "What IPTV means - and what it does not", new_page=False)
    add_body(doc, "Internet Protocol television means television delivered over IP networks. In a classic telecom deployment, live channels often traverse a managed access network as multicast MPEG transport streams, while on-demand material is unicast. In everyday usage, IPTV also refers to over-the-top live television delivered over the public Internet, most commonly as HTTP unicast and adaptive segments.", refs="S1")
    add_body(doc, "The delivery path can be centrally hosted or distributed through edge caches. A headend receives contribution feeds, encodes or passes through audio/video, creates service metadata, applies conditional access where required, and publishes the result to a managed network or CDN. A home server can perform a miniature version of the same job.")
    add_table(doc, ["Model", "Distribution", "Control", "Typical client contract"], [
        ["Managed operator IPTV", "Multicast live TV inside an ISP network; unicast VOD", "Operator QoS, gateway, set-top box", "Provisioned device/session; often closed"],
        ["OTT live streaming", "HTTP unicast through CDN", "App account, token, DRM/license", "HLS or DASH plus application APIs"],
        ["Open/custom IPTV", "Direct HTTP, HLS, RTP, RTSP, or local gateway", "M3U/XMLTV and local policy", "General-purpose player or TV server"],
        ["Network-tuner TV", "RF broadcast converted to LAN unicast/multicast", "Tuner discovery and reservation", "Lineup endpoint or backend API"],
        ["Virtual linear channel", "On-demand assets scheduled as a continuous stream", "Playout engine + generated EPG", "M3U/XMLTV or tuner emulation"],
    ], [1800, 2350, 2200, 3010], caption="Table 2. Common systems called IPTV.", small=True)
    add_callout(doc, "Terminology trap", "A file ending in .m3u8 may be a UTF-8 channel lineup or an HLS playlist. A lineup lists channels; an HLS master lists renditions; an HLS media playlist lists time-ordered segments. Inspect the tags and referenced URLs before assuming what it is.", refs="S2, S10")

    chapter(doc, "2", "Sources and IPTV provider models")
    add_body(doc, "A provider integration is a contract around four things: catalog, authorization, stream acquisition, and guide data. Some providers expose all four through one playlist URL. Others expose a lineup and guide separately. App-first services may expose only authenticated, DRM-protected playback through their own application and therefore cannot be represented faithfully by a static M3U file.")
    add_table(doc, ["Source type", "What you receive", "Operational concerns"], [
        ["Local broadcast tuner", "Antenna/cable/satellite services and in-band metadata", "Reception quality, tuner count, multiplex sharing, rescans"],
        ["Paid M3U/XMLTV service", "Lineup URL, stream URLs, guide URL", "Connection cap, token lifetime, geo/routing, terms and rights"],
        ["Managed operator feed", "Multicast groups or provisioned gateway service", "VLAN, IGMP, operator authentication, private routing"],
        ["FAST/public stream", "Public HLS/HTTP URL and sometimes no EPG", "URL churn, regional variants, ad markers, sparse metadata"],
        ["App-only commercial service", "App manifest, account session, DRM licenses", "No general-purpose M3U; CDM and device policy required"],
        ["Virtual channel engine", "Generated M3U/tuner API plus XMLTV", "Playout continuity, duration accuracy, transcode normalization"],
    ], [2050, 3250, 4060], caption="Table 3. Provider/source integration patterns.")
    add_body(doc, "Provider URLs commonly carry usernames, passwords, bearer tokens, signed query strings, cookies, or device identifiers. Treat a playlist as a secret. Do not commit it, paste it into issue trackers, expose it through an unauthenticated reverse proxy, or let raw URLs enter logs. Stable local aliases can shield downstream clients from token rotation.")
    add_body(doc, "Concurrency is a contractual and technical limit. One viewer, one recorder, and one probe can count as three open streams unless a proxy shares an upstream connection. Define exactly what consumes a slot and reserve capacity for scheduled recordings.")
    add_callout(doc, "Rights boundary", "Use streams and guide data for which you have permission. Publicly reachable is not the same as licensed for redistribution. DRM removal, credential sharing, and restreaming can violate contracts or law. Projects such as iptv-org maintain removal/blocklist processes, but operators remain responsible for their own use.", fill="FBEDED", accent=RED, refs="S36, S37, S41")

    chapter(doc, "3", "Network tuners: bridging RF television onto IP")
    add_body(doc, "A network tuner converts a physical broadcast input into an IP-accessible service. The tuner locks a frequency and modulation scheme; a demultiplexer selects one program from the received transport multiplex; a server or client reads the resulting MPEG-TS. Discovery and lineup APIs hide most RF details from media servers.")
    add_body(doc, "HDHomeRun-class devices expose device discovery data including tuner count and a lineup URL. Current models can publish virtual-channel lineup data at /lineup.json and stream a virtual channel over HTTP, while lower-level control can direct a tuner to a UDP target. The device is both a scarce resource pool and a source of truth about receivable channels.", refs="S11, S12")
    add_code(doc, """
http://TUNER_IP/discover.json
http://TUNER_IP/lineup.json
http://TUNER_IP:5004/auto/v9.1

# Lower-level example from a trusted LAN administration host:
hdhomerun_config discover
hdhomerun_config DEVICE_ID get /tuner0/streaminfo
hdhomerun_config DEVICE_ID set /tuner0/target udp://PLAYER_IP:5000
""", "HDHomeRun discovery and stream shapes")
    add_body(doc, "SAT>IP is another remote-tuner pattern, using discovery plus RTSP/HTTP to request services from satellite or other DVB frontends. Tvheadend, NextPVR, and similar backends generalize this further: they own adapters, networks, muxes, services, channels, EPG ingestion, and DVR rules, then expose normalized streams and metadata to clients.")
    add_table(doc, ["Resource", "Constraint", "Design implication"], [
        ["Frontend/tuner", "Can lock one RF tuning configuration at a time", "Overlapping channels may require multiple tuners"],
        ["Multiplex", "Carries several services on one frequency", "A capable backend may share one tuned mux among consumers"],
        ["LAN path", "Sustains the aggregate transport bitrate", "Use wired Ethernet where possible; isolate multicast intentionally"],
        ["Backend profile", "May pass through, remux, or transcode", "Browser clients often need a compatible profile"],
        ["Guide source", "OTA guide may be shallow or sparse", "Blend or replace with licensed XMLTV where appropriate"],
    ], [1800, 3200, 4360], caption="Table 4. Tuner-side resources and constraints.")
    add_body(doc, "For IPv4 multicast, hosts signal group interest through IGMP. Snooping switches use those reports to avoid flooding every LAN port, but a misconfigured querier, snooping bridge, VLAN boundary, or Wi-Fi multicast conversion can make a stream work briefly and then disappear. IGMPv3 now lives in RFC 9776; RFC 4541 remains useful switch guidance.", refs="S5, S6")

    chapter(doc, "4", "Lineups and extended M3U")
    add_body(doc, "A lineup answers: which channels exist, how should they be presented, and where can each be opened? Extended M3U is the most common interchange format. It is line-oriented and forgiving, which makes it portable but also creates dialects.")
    add_code(doc, """
#EXTM3U x-tvg-url="https://guide.example/denver.xml.gz"
#EXTINF:-1 tvg-id="kusa.denver.example" tvg-name="KUSA" \\
  tvg-chno="9.1" tvg-logo="https://img.example/kusa.png" \\
  group-title="Local;News",KUSA 9.1
https://edge.example/live/kusa/master.m3u8?token=REDACTED
""", "A portable channel-lineup entry")
    add_table(doc, ["Field", "Role", "Stability guidance"], [
        ["tvg-id", "Join key to XMLTV channel@id by convention", "Stable and unique within the guide namespace"],
        ["tvg-name", "Human/matcher hint", "Useful fallback; not a canonical key"],
        ["tvg-chno", "Display/sort number", "Can change by region or user preference"],
        ["tvg-logo", "Presentation artwork", "Cacheable; do not use as identity"],
        ["group-title", "One or more categories/groups", "A snapshot; syntax and multi-group support vary"],
        ["URL", "Current acquisition route", "Expect tokens, hosts, or fallback order to change"],
        ["#EXTVLCOPT / #KODIPROP", "Player-specific options", "Treat as nonportable extensions"],
    ], [1750, 3600, 4010], caption="Table 5. Common IPTV M3U fields and their proper role.", small=True)
    add_body(doc, "Kodi's IPTV Simple documentation illustrates the real ecosystem: tvg-id, tvg-name, tvg-chno, tvg-logo, group-title, radio, provider metadata, catch-up modes, and player-specific properties. It also shows that some tags are de facto conventions, not one universal M3U standard. Validate against the actual consumer.", refs="S10")
    add_body(doc, "Lineup ingestion should be deterministic: fetch with a timeout, verify HTTP status and content type, decode UTF-8 safely, parse into a temporary structure, reject or quarantine malformed entries, resolve duplicates using policy, then atomically replace the active snapshot. Retain the last known good lineup when refresh fails.")
    add_callout(doc, "Do not key by URL", "A signed URL may rotate every hour while the logical channel remains unchanged. Store a stable canonical channel ID and versioned route records with health, priority, and expiry metadata.", fill=LIGHT_GOLD, accent=GOLD)

    chapter(doc, "5", "XMLTV, EPG data, and time")
    add_body(doc, "An electronic program guide is a time-indexed catalog of airings. XMLTV represents it as channel records followed by programme records. A programme's channel attribute must refer to a channel id; its required start and channel attributes say when and where the airing begins. Stop, title, subtitle, description, categories, credits, episode numbering, ratings, images, language, audio/video properties, and repeat status add scheduling and presentation value.", refs="S8, S9")
    add_code(doc, """
<?xml version="1.0" encoding="UTF-8"?>
<tv source-info-name="Example Guide" generator-info-name="guide-builder/1.0">
  <channel id="kusa.denver.example">
    <display-name>KUSA 9.1</display-name>
    <icon src="https://img.example/kusa.png"/>
  </channel>
  <programme start="20260820000000 +0000"
             stop="20260820003000 +0000"
             channel="kusa.denver.example">
    <title lang="en">Evening News</title>
    <category lang="en">News</category>
    <episode-num system="onscreen">2026-08-19</episode-num>
    <new/>
  </programme>
</tv>
""", "Minimal useful XMLTV")
    add_body(doc, "XMLTV time strings are compact ISO-8601-like values. The DTD permits partial precision and says UTC is assumed when no timezone is present, but production systems should emit full timestamps with an explicit numeric offset or normalize to +0000. That removes ambiguity at daylight-saving transitions. Treat programme intervals as half-open: start is included and stop is excluded, so back-to-back programmes touch without overlapping.", refs="S8")
    add_table(doc, ["Problem", "Symptom", "Control"], [
        ["Missing/unstable channel ID", "No guide or wrong guide after refresh", "Canonical mapping table; never fuzzy-match silently"],
        ["Timezone/DST error", "Guide shifted by one or more hours", "Explicit offsets; test both DST transitions"],
        ["Sparse stop times", "Incorrect grid widths or recording ends", "Infer only under controlled policy; flag uncertainty"],
        ["Duplicate overlapping events", "Stacked/hidden guide rows", "Source priority plus overlap validation"],
        ["Episode numbering mismatch", "Series rules or metadata fail", "Preserve onscreen and provider IDs; understand xmltv_ns is zero-based"],
        ["Stale feed", "Recordings follow yesterday's schedule", "Freshness watermark, coverage horizon, last-known-good fallback"],
    ], [2200, 3300, 3860], caption="Table 6. EPG data-quality failures.")
    add_body(doc, "OTA systems carry guide metadata in-band: ATSC PSIP defines virtual-channel and event tables; DVB Service Information uses service and Event Information Tables, including present/following and schedule forms. A backend can ingest these directly, translate them to an internal database, and export XMLTV for downstream systems.", refs="S39, S40, S42")
    add_body(doc, "A guide refresh job should measure coverage, not merely parse success: percentage of active channels mapped, percentage with a current event, hours of future coverage, overlap/gap counts, oldest source timestamp, and guide age at publish time.")

    chapter(doc, "6", "Channel IDs and the identity join")
    add_figure(doc, ASSET_DIR / "identity.png", "Figure 2. Channel identity crosses several namespaces.", "Six identity layers from physical RF through MPEG transport, M3U, XMLTV, internal application IDs, and user-facing presentation.")
    add_body(doc, "The XMLTV DTD requires each channel id to be unique and recommends a DNS-like form. Programme records refer to that id. In an IPTV M3U, tvg-id is the de facto join key to the XMLTV channel id. This convention is extremely important, but tvg-id itself is not defined by the XMLTV DTD.", refs="S8, S10")
    add_table(doc, ["Identifier", "Namespace", "Example", "Safe use"], [
        ["Canonical channel ID", "Your registry", "us.co.denver.kusa.main", "Primary persistent key"],
        ["XMLTV channel@id", "Guide provider", "kusa.denver.example", "Join programmes to channel"],
        ["M3U tvg-id", "Playlist/provider", "kusa.denver.example", "Join lineup to guide"],
        ["Service/program ID", "MPEG-TS multiplex", "service_id 0x0003", "Select program within transport"],
        ["Channel number", "Lineup/user", "9.1 or 609", "Display and remote-control entry"],
        ["Internal server ID", "Jellyfin/backend DB", "UUID/opaque key", "Persist schedules, permissions, favorites"],
        ["Stream URL", "Provider/CDN", "signed HLS URL", "Mutable route, never identity"],
    ], [1980, 1780, 2140, 3460], caption="Table 7. Identity fields should not be conflated.", small=True)
    add_body(doc, "A robust mapper uses exact configured IDs first. Optional fallbacks can compare normalized callsigns, names, channel numbers, region, language, and logo host, but ambiguous matches should be presented for review. Persist the operator's choice so later refreshes do not undo it.")
    add_code(doc, """
channel_id: us.co.denver.kusa.main
display_name: KUSA 9.1
number: "9.1"
groups: [Local, News]
identifiers:
  xmltv: kusa.denver.example
  m3u_tvg_id: kusa.denver.example
  atsc_source_id: 1234
routes:
  - kind: hdhomerun
    uri: http://tuner.local:5004/auto/v9.1
    priority: 10
  - kind: hls
    uri_secret_ref: provider/kusa
    priority: 20
""", "A canonical channel registry record")
    add_callout(doc, "Duplicate services", "SD, HD, 4K, east/west, local-insert, and language feeds may share branding but are distinct routable services. Decide whether they share guide identity, programme identity, neither, or both - then encode that decision explicitly.", fill=LIGHT_TEAL, accent=TEAL)

    chapter(doc, "7", "Scheduling, recording, timeshift, and catch-up")
    add_body(doc, "A scheduler converts guide events and user rules into resource reservations. It must account for real start and stop times, pre/post padding, tuner/provider limits, source priority, storage, and post-processing. Tvheadend exposes event timers, time-based timers, autorecord/series rules, and DVR profiles; its API distinguishes scheduled times from real padded times.", refs="S26, S27")
    add_step(doc, "Resolve a rule to candidate airings using stable programme metadata: series identifier when available, then title/subtitle/category policy.", decimal_id)
    add_step(doc, "Apply eligibility policy: new-only, channel/group restrictions, time window, quality preference, repeat detection, retention.", decimal_id)
    add_step(doc, "Expand each airing by pre-roll, post-roll, and device warm-up; this is the real reservation interval.", decimal_id)
    add_step(doc, "Allocate a tuner or provider slot across all overlapping reservations. If multiplex sharing exists, model it explicitly rather than assuming it.", decimal_id)
    add_step(doc, "Choose a route and recording profile, reserve storage, start early, validate bytes/timestamps, and record atomically to a temporary name.", decimal_id)
    add_step(doc, "Finalize metadata, run optional post-processing, verify duration/streams, publish the recording, and record failure reason if incomplete.", decimal_id, refs="S26, S27")
    add_body(doc, "Padding improves resilience to schedule drift but increases conflicts. The resource model must use padded intervals. Two adjacent recordings on one tuner overlap if the first has five minutes of post-roll and the second has two minutes of pre-roll, even when the published programmes do not overlap.")
    add_table(doc, ["Input bitrate", "Approx. per hour", "24 hours continuous", "30 days continuous"], [
        ["4 Mb/s", "1.8 GB", "43 GB", "1.3 TB"],
        ["8 Mb/s", "3.6 GB", "86 GB", "2.6 TB"],
        ["12 Mb/s", "5.4 GB", "130 GB", "3.9 TB"],
        ["20 Mb/s", "9.0 GB", "216 GB", "6.5 TB"],
        ["30 Mb/s", "13.5 GB", "324 GB", "9.7 TB"],
    ], [1800, 2200, 2500, 2860], caption="Table 8. Rule-of-thumb recording storage (decimal GB/TB, excluding filesystem overhead).")
    add_body(doc, "Timeshift is a local rolling buffer that lets a viewer pause or seek within recently received live content. Catch-up is provider-side archival playback, often constructed from programme start/duration and provider-specific URL rules. They look similar in the UI but have different owners, retention, failure modes, and rights. Kodi's IPTV Simple examples show several nonportable catch-up URL conventions.", refs="S10")
    add_callout(doc, "Clock discipline", "Schedules use civil time; media uses stream clocks. Keep the server synchronized with NTP, normalize EPG timestamps, preserve source PTS/PCR where sound, and record both planned and observed start/stop times.", fill=LIGHT_GOLD, accent=GOLD)

    chapter(doc, "8", "Transport and streaming protocols")
    add_body(doc, "The transport choice trades latency, scalability, recoverability, firewall traversal, cacheability, and player reach. Protocol and container are separate: MPEG-TS can travel over UDP, RTP, HTTP, HLS segments, or SRT; H.264 can be carried in TS or fMP4.")
    add_table(doc, ["Protocol", "Shape", "Strength", "Typical IPTV use"], [
        ["UDP", "Push datagrams; unicast or multicast", "Low overhead and latency", "Managed LAN/headend transport; loss is visible"],
        ["RTP/RTCP", "Timed media packets plus control reports", "Sequence/timestamp/jitter semantics", "Real-time unicast/multicast and RTSP media"],
        ["RTSP", "Control session, usually with RTP media", "DESCRIBE/SETUP/PLAY/PAUSE", "Cameras, SAT>IP, remote tuners"],
        ["HTTP continuous TS", "One long pull response", "Simple firewall traversal", "Direct tuner/provider streams"],
        ["HLS", "Master/media playlists plus segments", "CDN-friendly, adaptive, broad reach", "OTT live TV and server-to-browser delivery"],
        ["MPEG-DASH", "MPD plus representations/segments", "Codec/container flexibility and ABR", "OTT applications and browsers"],
        ["SRT", "Reliable, encrypted UDP transport", "ARQ and configurable latency", "Contribution and backhaul over imperfect links"],
        ["WebRTC", "RTP-centered interactive stack", "Very low latency, congestion control", "Interactive or real-time viewing, not a lineup format"],
    ], [1450, 2400, 2500, 3010], caption="Table 9. Transport choices at a glance.", small=True)
    add_body(doc, "RTP sequence numbers reveal loss and order; RTP timestamps reconstruct sampling time; RTCP reports carry reception statistics such as packet loss and interarrival jitter. RTSP is the control plane around a presentation and transport selection, not the media container itself.", refs="S3, S4")
    add_body(doc, "HLS uses a master playlist to describe variant streams and a media playlist to enumerate time-ordered segments. Segments may be MPEG-TS or fragmented MP4; each fMP4 media playlist uses an initialization map. The live edge moves as media sequence numbers advance. Latency is driven by encoder delay, GOP length, segment/part duration, playlist window, publication cadence, network transfer, and player buffer.", refs="S2")
    add_code(doc, """
#EXTM3U
#EXT-X-VERSION:7
#EXT-X-TARGETDURATION:4
#EXT-X-MEDIA-SEQUENCE:8124
#EXT-X-MAP:URI="init.mp4"
#EXTINF:4.000,
segment-8124.m4s
#EXTINF:4.000,
segment-8125.m4s
""", "An HLS media playlist (not a channel lineup)")
    add_body(doc, "Adaptive bitrate players choose among representations using recent throughput, buffer occupancy, and device/display constraints. ABR is therefore a feedback controller, not a promise to always pick the highest bitrate.", refs="S7")
    add_body(doc, "SRT wraps payload-agnostic data in a loss-recovering transport with a configured latency budget and optional AES encryption. ARQ needs enough time for a loss report and retransmission; setting latency below the path's recovery needs converts recoverable loss into visible damage.", refs="S38")

    chapter(doc, "9", "Containers, codecs, tracks, and timestamps")
    add_callout(doc, "Four layers", "Playlist/manifest tells a client what to request. Transport moves bytes. Container/multiplex organizes timed tracks. Codec compresses a track. Saying a stream is 'H.264' does not tell you whether a browser can open its protocol, container, audio, or subtitles.", fill=LIGHT_TEAL, accent=TEAL)
    add_body(doc, "MPEG transport stream is built from 188-byte packets. Packet identifiers (PIDs) separate tables and elementary streams. The Program Association Table points to Program Map Tables; a PMT lists the audio, video, subtitle/data tracks for a program. Program Clock Reference anchors the decoder clock; PTS says when to present a frame/sample; DTS says when decoding must occur when reordering is involved.", refs="S32")
    add_table(doc, ["Layer", "Common choices", "Important compatibility dimensions"], [
        ["Video codec", "MPEG-2, H.264/AVC, HEVC/H.265, VP9, AV1", "Profile, level, bit depth, chroma, interlace, frame rate, hardware decoder"],
        ["Audio codec", "MP2, AAC, AC-3, E-AC-3, Opus", "Channels/layout, sample rate, passthrough, browser support"],
        ["Subtitles/captions", "CEA-608/708, DVB bitmap, teletext, WebVTT, TTML, SRT/ASS", "In-band vs sidecar; text vs image; burn-in requirement"],
        ["Container", "MPEG-TS, fMP4/MP4, Matroska, WebM", "Track types, timestamp model, streamability, browser demux support"],
        ["Color/HDR", "SDR BT.709, HDR10, HLG, Dolby Vision", "Transfer, primaries, metadata, tone mapping, display path"],
    ], [1900, 3160, 4300], caption="Table 10. Media compatibility is multidimensional.")
    add_body(doc, "H.264 remains the broadest common delivery codec; HEVC improves compression but has more uneven browser/device support; AV1 targets higher compression efficiency and modern high-resolution streaming. A nominal codec name is insufficient: 10-bit H.264, high HEVC levels, or unsupported reference structures may still force a transcode.", refs="S33, S34")
    add_body(doc, "Matroska can carry many track types and uses EBML structures, Tracks metadata, Clusters, and timestamp scales. It is excellent for storage, but browser clients may require remuxing to an HLS-friendly container even when the encoded audio/video is already compatible.", refs="S35, S15")
    add_body(doc, "Broadcast video may be interlaced. A client that cannot deinterlace well may show combing, while server-side deinterlacing forces decoded-frame processing and usually video re-encoding. Likewise, image-based subtitles often require burn-in, turning an otherwise cheap remux into a full transcode.")

    chapter(doc, "10", "How VLC opens and buffers an IPTV stream")
    add_body(doc, "VLC is a modular pipeline. An access module opens the URL/device; a demuxer identifies the container and separates elementary streams; packetizers/decoders turn compressed tracks into timed frames and samples; the clock system synchronizes audio, video, and subtitles; outputs render or send media onward. VLC's advertised inputs include UDP/RTP unicast and multicast, HTTP, RTP over TCP, DVB, MPEG-TS, Matroska, MP4, and many codecs.", refs="S22, S24")
    add_figure(doc, ASSET_DIR / "buffer_chain.png", "Figure 3. Playback buffering exists at several layers.", "Network, access, demux, decode and output buffering stages with a latency-versus-resilience warning.")
    add_body(doc, "The network-caching setting is expressed in milliseconds for network resources. VLC source also distinguishes file, disc, and live-capture caching and exposes clock-jitter and synchronization controls. Modules interpret the requested delay in their own path; for example, VLC's live555 RTSP access uses network-caching as the requested PTS delay. This is why one cache number does not behave identically for every protocol.", refs="S21, S22")
    add_code(doc, """
vlc --network-caching=1500 "https://edge.example/live/master.m3u8"
vlc --network-caching=1000 "http://tuner.local:5004/auto/v9.1"
vlc --network-caching=800  "udp://@239.10.20.30:5000"
vlc --network-caching=1000 "rtsp://camera.example/live"
""", "Opening representative network sources")
    add_body(doc, "A larger cache can absorb jitter, short stalls, or bursty segment downloads, but increases tune/start latency and distance from live. A smaller cache feels responsive but exposes variance. Neither fixes a source whose average throughput is below the encoded bitrate, a missing PAT/PMT, corrupted timestamps, a GOP with no reachable keyframe, or an unsupported codec.")
    add_table(doc, ["Observed VLC symptom", "Likely layer", "Next evidence"], [
        ["Opens slowly before tracks appear", "Probe/demux or HLS manifest", "Messages log; manifest; PAT/PMT; probe size"],
        ["Video starts only after several seconds", "Keyframe/GOP or cache target", "Frame types; keyframe interval; network cache"],
        ["Periodic rebuffering", "Throughput/jitter/segment availability", "Segment timings, bitrate, packet loss, CDN status"],
        ["Audio leads/lags", "PTS/PCR discontinuity or decoder path", "Verbose clock/timestamp logs; ffprobe packets"],
        ["Blockiness on UDP", "Packet loss or socket overrun", "Interface drops, IGMP, UDP receive buffer, RTP stats"],
        ["Works in VLC, fails in browser", "Browser protocol/container/codec limits", "Client capability matrix; Jellyfin play method"],
    ], [2500, 3000, 3860], caption="Table 11. VLC is a diagnostic endpoint as well as a player.")
    add_body(doc, "VLC can also act as a stream processor. Its stream-output chain composes modules such as transcode and standard output. For HTTP output, the client pulls the published stream; the mux must be declared, commonly MPEG-TS.", refs="S23")
    add_code(doc, """
vlc input.mp4 --sout="#std{access=http,mux=ts,dst=:8090/channel}"
""", "A small VLC HTTP-TS relay")

    chapter(doc, "11", "How FFmpeg inspects, remuxes, and transcodes")
    add_body(doc, "FFmpeg's pipeline is explicit: protocol I/O feeds a demuxer; compressed packets are either stream-copied or decoded; decoded frames may pass through filters; encoders create new compressed packets; a muxer writes the target container. Streamcopy (-c copy) skips decode/filter/encode, so it is fast and lossless, but it cannot fix codec incompatibility or apply frame-level filters.", refs="S17")
    add_code(doc, """
ffprobe -v error -show_format -show_programs -show_streams \\
  -of json "INPUT_URL"

ffprobe -v error -select_streams v:0 -show_packets \\
  -show_entries packet=pts_time,dts_time,duration_time,flags \\
  -read_intervals "%+10" -of csv=p=0 "INPUT_URL"
""", "Inspect before changing anything")
    add_body(doc, "ffprobe can report format, programs, streams, packets, and frames in JSON, XML, or compact text. For MPEG-TS, show_programs is valuable because one transport can contain multiple services. Probe/analyze limits trade startup time against the chance of discovering sparse tracks.", refs="S19, S20")
    add_code(doc, """
# Record the selected program/tracks without quality loss:
ffmpeg -hide_banner -loglevel warning -i "INPUT_URL" \\
  -map 0:v:0 -map 0:a:0? -map 0:s? -c copy -t 01:00:00 recording.ts

# Remux to another container when the codecs are legal there:
ffmpeg -i input.ts -map 0:v:0 -map 0:a:0? -c copy output.mkv
""", "Streamcopy and explicit mapping")
    add_body(doc, "Use -map rather than relying on automatic selection when channel streams may contain multiple languages, commentary tracks, data PIDs, or image subtitles. A trailing ? makes a map optional. The output container must support every copied codec and required metadata.", refs="S17")
    add_code(doc, """
# HTTP input with bounded waits and reconnect policy:
ffmpeg -rw_timeout 15000000 \\
  -reconnect 1 -reconnect_on_network_error 1 \\
  -reconnect_on_http_error 4xx,5xx -reconnect_streamed 1 \\
  -reconnect_delay_max 10 -i "HTTPS_INPUT" -map 0 -c copy output.ts

# UDP input with a larger circular buffer and nonfatal overruns:
ffmpeg -i "udp://239.10.20.30:5000?fifo_size=1000000&overrun_nonfatal=1&timeout=5000000" \\
  -map 0 -c copy output.ts
""", "Network resilience controls")
    add_body(doc, "FFmpeg documents rw_timeout for network I/O, HTTP reconnect policies, and UDP socket/circular buffers. Reconnect can preserve a long-running job, but it can also duplicate, gap, or reset timestamps depending on the source. The recording layer must detect discontinuities rather than assuming a reconnect is seamless.", refs="S18")
    add_code(doc, """
# Generate a six-segment sliding HLS window by streamcopy:
ffmpeg -i input.ts -map 0:v:0 -map 0:a:0? -c copy \\
  -f hls -hls_time 4 -hls_list_size 6 \\
  -hls_flags delete_segments+temp_file output.m3u8

# Normalize video/audio for a broadly compatible 30 fps HLS rendition:
ffmpeg -i input.ts -map 0:v:0 -map 0:a:0? \\
  -c:v libx264 -preset veryfast -g 120 -keyint_min 120 -sc_threshold 0 \\
  -c:a aac -b:a 128k -f hls -hls_time 4 -hls_list_size 6 \\
  -hls_flags delete_segments+temp_file output.m3u8
""", "HLS remux versus normalization")
    add_body(doc, "The HLS muxer cuts at the next keyframe after the target duration. With streamcopy, FFmpeg cannot invent better keyframe placement; long or irregular GOPs produce long/irregular segments. When encoding, align a closed GOP to the segment duration (for example, 4 seconds × 30 fps = 120 frames) and disable scene-cut keyframes if strict alignment matters.", refs="S19")
    add_body(doc, "Timestamp flags are surgical tools. -fflags +genpts may generate missing PTS; -copyts preserves input timestamps; -start_at_zero shifts preserved timestamps; avoid_negative_ts can move output timestamps. They interact with demuxer and muxer behavior. Capture evidence first and verify A/V sync after any repair.", refs="S17")
    add_callout(doc, "Remuxing is not transcoding", "Remux: packets are copied into a new container or segment layout. Transcode: packets are decoded and new packets are encoded. Remuxing is usually low-CPU and quality-neutral; transcoding costs compute, adds latency, and changes quality.", fill=LIGHT_BLUE, accent=BLUE)

    chapter(doc, "12", "How Jellyfin handles live streams")
    add_body(doc, "Jellyfin's Live TV model separates tuner devices from guide providers. It supports HDHomeRun and M3U tuners directly, with additional backends through plugins. An M3U tuner can point to a local or HTTP playlist, specify a user agent, enforce a simultaneous-stream limit, and optionally auto-loop problematic live streams. Guide data is then added and mapped to physical/tuner channels.", refs="S13")
    add_body(doc, "Current Jellyfin documentation says a Live TV configuration chooses Schedules Direct or XMLTV rather than using both simultaneously in the built-in flow. Channel mapping is explicit. Jellyfin can also consume Tvheadend's M3U and XMLTV endpoints, although the dedicated plugin is recommended where available.", refs="S13")
    add_figure(doc, ASSET_DIR / "playback_decision.png", "Figure 4. Jellyfin selects the least expensive compatible playback path.", "Decision from source compatibility to Direct Play, Remux, Direct Stream audio transcode, or full video transcode.")
    add_body(doc, "The client sends capability profiles describing codecs, containers, resolution, bitrate, and other constraints. Jellyfin then selects Direct Play, Remux, Direct Stream, or Transcode. In Jellyfin terminology, Direct Stream typically means video copy with audio transcoding; a video codec mismatch forces video transcoding. Subtitles can trigger remuxing or burn-in.", refs="S14, S15")
    add_table(doc, ["Play method", "What changes", "Server cost", "Typical trigger"], [
        ["Direct Play", "Nothing in media", "Minimal", "Client accepts protocol/container/all selected tracks"],
        ["Remux", "Container and/or segmentation", "Low", "Codec tracks supported, container not supported"],
        ["Direct Stream", "Audio and/or subtitle representation", "Moderate", "Video supported; audio or text path is not"],
        ["Transcode", "Video (often audio too)", "High", "Video codec/profile/bitrate/resolution/subtitle burn-in"],
    ], [1800, 2700, 1500, 3360], caption="Table 12. Jellyfin playback modes.")
    add_body(doc, "Jellyfin uses its maintained jellyfin-ffmpeg build for media processing and can use Intel, AMD, Nvidia, Apple, and other hardware paths where supported. Hardware decode and encode capabilities are not symmetric; a GPU may decode a source format but not encode the required target, or tone mapping may introduce another constraint.", refs="S16")
    doc.add_paragraph("Jellyfin live-stream triage", style="Heading 2")
    jellyfin_steps_id = add_numbering(doc, "decimal", "%1.", 540, 270)
    add_step(doc, "Open Dashboard during playback and record the listed play method and reason for transcoding.", jellyfin_steps_id)
    add_step(doc, "Inspect the corresponding FFmpeg log: input probe, selected streams, filters, encoders, HLS output, warnings, and exit code.", jellyfin_steps_id)
    add_step(doc, "Compare the exact client capability profile with the source: protocol, container, video profile/level/bit depth, audio layout, subtitles, and bitrate.", jellyfin_steps_id)
    add_step(doc, "Test the tuner/URL independently with ffprobe and VLC from the Jellyfin host, not only from a desktop on another network path.", jellyfin_steps_id)
    add_step(doc, "Verify transcode storage, permissions, free space, hardware device access, and processing speed above 1.0x for sustained live playback.", jellyfin_steps_id)
    add_step(doc, "Only then change cache, probe, or hardware settings, one variable at a time, and keep the before/after log.", jellyfin_steps_id)
    add_callout(doc, "Container surprise", "Jellyfin may remux a source to HLS/TS for a web client even when the video codec is supported. This is expected: client compatibility includes the delivery and container path, not just H.264 versus HEVC.", refs="S15")

    chapter(doc, "13", "Dynamic groups, dynamic channels, and proxies")
    add_body(doc, "An M3U group is merely metadata in a lineup snapshot. A dynamic group is created when a generator evaluates rules - country, language, genre, source, resolution, entitlement, health, favorites, or current event type - and republishes membership. The downstream player still receives an ordinary group-title or application-specific collection.")
    add_body(doc, "M3U/XMLTV proxies such as Threadfin can merge sources, filter channels, map guide IDs, order and renumber channels, assign logos/categories, enforce tuner limits, re-stream, and provide backup routes. This creates a stable local contract for Jellyfin/Plex/Emby while upstream feeds change.", refs="S25")
    add_table(doc, ["Dynamic behavior", "Implementation pattern", "Identity rule"], [
        ["Rule-based group", "Evaluate registry/EPG/health fields during publish", "Membership changes; channel ID does not"],
        ["Ephemeral event channel", "Create before event; retire after archive window", "Use stable event-service ID; avoid reusing another channel's ID"],
        ["Failover route", "Multiple prioritized URLs behind one local channel URL", "Route changes; channel/guide ID stays fixed"],
        ["Quality variant", "Expose variants separately or policy-select one", "Separate service IDs if users can choose them"],
        ["Virtual linear channel", "Schedule library assets and emit continuous stream + XMLTV", "Channel ID stable; programmes generated from playout"],
        ["On-demand pseudo-channel", "Schedule advances only while watched", "Guide requires explicit placeholder/update semantics"],
    ], [2300, 3650, 3410], caption="Table 13. Dynamic channel patterns.")
    add_body(doc, "Virtual-channel engines schedule on-demand assets into a linear timeline. ErsatzTV uses collections, smart collections that update from searches/lists, and playout instructions for sequencing, padding, filler, loops, and EPG grouping. dizqueTV can expose generated channels through M3U or HDHomeRun emulation and writes XMLTV guide data.", refs="S28, S29, S30")
    add_body(doc, "Publish lineup and guide snapshots as a coordinated release. Generate both to temporary paths, validate cross-references and guide horizon, assign a release version, then switch stable URLs atomically. A lineup referring to a new tvg-id before its XMLTV channel appears produces a temporary blank guide; the inverse leaves orphan guide records.")
    add_callout(doc, "Cache invalidation", "Downstream servers may refresh M3U and XMLTV on different schedules. Use stable endpoints, ETag/Last-Modified where useful, bounded cache times, a monotonically increasing release ID, and an operator-triggered refresh path.", fill=LIGHT_GOLD, accent=GOLD)

    chapter(doc, "14", "Reliability, networking, and observability")
    add_body(doc, "A reliable IPTV service validates each stage independently. 'The URL returns 200' is not a media health check. A good probe verifies connection, sustained bytes, container recognition, service tables, selected audio/video tracks, timestamp progress, decodable keyframes, and an acceptable error rate.")
    add_table(doc, ["Metric", "Why it matters", "Example alert"], [
        ["Lineup/guide fetch age", "Detects stale control-plane data", "No successful refresh within 2 planned intervals"],
        ["Mapped-channel ratio", "Measures EPG join completeness", "Below 98% of enabled channels"],
        ["Current-event coverage", "Finds empty/stale guide rows", "No current programme for >5% of active channels"],
        ["Tune startup p50/p95", "Captures route, probe, keyframe, and buffer delay", "p95 doubles from baseline"],
        ["Continuity/timestamp errors", "Predicts blockiness and A/V drift", "Nonzero sustained rate or burst above threshold"],
        ["Upstream slots/tuners used", "Prevents recording conflicts", "Utilization above reserved capacity"],
        ["Transcode speed and queue", "Real time needs speed >1x with margin", "Below 1.1x for live or jobs queued"],
        ["Recording validation", "Catches zero-byte/short/wrong-track files", "Duration or byte rate below policy"],
    ], [2300, 4100, 2960], caption="Table 14. Operational signals for an IPTV platform.", small=True)
    add_body(doc, "Size links by aggregate bitrate plus overhead and headroom. Four simultaneous 20 Mb/s channels already consume 80 Mb/s before protocol overhead, retransmissions, other traffic, or server-to-client duplication. Wi-Fi may advertise high PHY rates while delivering unstable multicast or sustained throughput; wired Ethernet is the conservative server/tuner path.")
    add_body(doc, "For multicast, inspect the entire path: source interface, VLAN, querier, router/proxy, switch snooping state, wireless bridge, host join, socket receive drops, and firewall. For HTTP/HLS, inspect DNS, TCP/TLS setup, redirects, authentication, CDN edge, manifest age, segment availability, transfer time, retries, and origin 4xx/5xx.")
    add_body(doc, "Backoff and circuit breaking prevent a dead channel from becoming a self-inflicted denial of service. Separate fast viewer tune attempts from slower background health probes. Maintain last-known-good metadata even when media is unavailable, and show a precise failure state instead of silently substituting the wrong channel.")
    add_callout(doc, "Capacity reserve", "Do not let background probes consume every provider slot or tuner. Budget viewers, scheduled recordings, warm-up overlap, retries, and at least one operational reserve where the source allows it.", fill="FBEDED", accent=RED)

    chapter(doc, "15", "Security, privacy, and lawful operation")
    add_bullet(doc, "Store provider credentials and signed playlist URLs in a secret manager or protected configuration; expose redacted local aliases to logs and clients.", bullet_id)
    add_bullet(doc, "Use TLS for remote lineup, guide, segment, API, and administration traffic; validate certificates rather than disabling verification.", bullet_id)
    add_bullet(doc, "Authenticate local M3U/XMLTV/proxy endpoints when they reveal entitlements, topology, viewing options, or credentials.", bullet_id)
    add_bullet(doc, "Constrain parsers: maximum download size, decompression limits, XML external-entity protections, URL allowlists where appropriate, and timeouts.", bullet_id)
    add_bullet(doc, "Run FFmpeg/VLC post-processors with least privilege; treat media and metadata as untrusted input; keep packages patched.", bullet_id)
    add_bullet(doc, "Separate admin, playback, and recording permissions. Jellyfin exposes Live TV access and recording-management controls per user.", bullet_id, refs="S13")
    add_bullet(doc, "Document content rights, territories, household/device limits, recording/retention permissions, and redistribution prohibitions before integrating a source.", bullet_id)
    add_body(doc, "Commercial OTT commonly combines authentication, encrypted segments, and a license exchange with a Content Decryption Module. W3C EME standardizes browser interaction with key systems; it does not define a DRM system or grant rights. FairPlay is one protected-HLS implementation on Apple platforms. A static M3U cannot reproduce these application and license policies by itself.", refs="S36, S37")
    add_callout(doc, "Privacy", "Guide queries, channel-change logs, playback history, and recordings can reveal viewing habits. Minimize retention, restrict access, avoid embedding account IDs in exported URLs, and scrub diagnostic bundles before sharing them.", fill=LIGHT_TEAL, accent=TEAL)

    chapter(doc, "16", "A reference architecture for a maintainable home or small-service stack")
    add_body(doc, "The simplest maintainable design creates one local contract for downstream consumers. It is not necessary to deploy every component; each layer exists to isolate a kind of change.")
    architecture_steps_id = add_numbering(doc, "decimal", "%1.", 540, 270)
    add_step(doc, "Acquire sources: network tuner/backend, authorized provider feeds, public streams, or virtual playout.", architecture_steps_id)
    add_step(doc, "Normalize into a canonical channel registry with stable IDs, mutable routes, display metadata, groups, provider limits, and guide identifiers.", architecture_steps_id)
    add_step(doc, "Fetch and validate XMLTV/OTA guide data; transform time and metadata; map explicitly to canonical channels.", architecture_steps_id)
    add_step(doc, "Publish versioned, atomic local M3U and XMLTV endpoints. A proxy may handle filtering, renumbering, route failover, and tuner emulation.", architecture_steps_id)
    add_step(doc, "Configure the TV server (for example Jellyfin) against those stable endpoints; set conservative stream limits and guide-refresh cadence.", architecture_steps_id)
    add_step(doc, "Prefer direct play or remux. Add FFmpeg normalization profiles only for demonstrated client incompatibilities.", architecture_steps_id)
    add_step(doc, "Run health probes, capacity checks, and recording validation outside the viewer request path; reserve source capacity.", architecture_steps_id)
    add_step(doc, "Back up the registry, mapping decisions, schedules, server configuration, and virtual-channel definitions - not just recordings.", architecture_steps_id)
    add_table(doc, ["Component", "Owns", "Should not own"], [
        ["Source adapter", "Authentication, discovery, stream open, source limits", "User-facing canonical identity"],
        ["Registry/generator", "Stable identity, groups, route policy, M3U/XMLTV publish", "Heavy media transformation"],
        ["Guide pipeline", "XMLTV/OTA ingest, time normalization, mappings, coverage", "Stream availability"],
        ["Proxy/backend", "Tuner allocation, failover, re-stream, DVR interface", "Provider secrets in exported URLs"],
        ["Jellyfin/media server", "Users, client negotiation, DVR UX, delivery/transcode", "Guessing ambiguous EPG matches"],
        ["Player", "Buffer, demux/decode, track selection, rendering", "Long-term provider/catalog normalization"],
    ], [2050, 3650, 3660], caption="Table 15. Clean ownership boundaries reduce coupling.")
    add_callout(doc, "Good default", "Normalize metadata eagerly; transform media lazily. It is cheap to make IDs, groups, and guide data consistent. It is expensive to transcode every channel 'just in case.'", fill=LIGHT_TEAL, accent=TEAL)

    chapter(doc, "17", "Troubleshooting runbook")
    add_body(doc, "Work from source toward client and stop at the first failing boundary. Preserve the original URL privately, timestamps, host, client, and exact command/log. Do not change three buffer/transcode/network settings at once.")
    add_table(doc, ["Stage", "Question", "Minimal test"], [
        ["1. Catalog", "Is the channel present and uniquely identified?", "Inspect parsed M3U and canonical registry"],
        ["2. Guide", "Does tvg-id map exactly to XMLTV channel@id?", "Count mapped/current/future programmes"],
        ["3. Authorization", "Can the server host open the route?", "HTTP status/redirect/TLS without exposing tokens"],
        ["4. Transport", "Do bytes arrive continuously?", "Network counters, packet capture where authorized"],
        ["5. Demux", "Are container, program, and tracks recognized?", "ffprobe show_format/show_programs/show_streams"],
        ["6. Timing/decode", "Do timestamps progress and frames decode?", "Short ffmpeg null decode and packet sample"],
        ["7. Server", "Which playback method and why?", "Jellyfin dashboard + FFmpeg log"],
        ["8. Client", "Can this exact client render the chosen output?", "Compare another client; capability profile"],
    ], [1600, 3600, 4160], caption="Table 16. A source-to-screen diagnostic ladder.")
    add_code(doc, """
# 15-second decode test; writes no media file:
ffmpeg -v warning -t 15 -i "INPUT_URL" -map 0:v:0 -map 0:a:0? -f null -

# Short lossless capture for private analysis:
ffmpeg -v warning -t 30 -i "INPUT_URL" -map 0 -c copy sample.ts

# Inspect the capture without network variability:
ffprobe -v error -show_format -show_programs -show_streams -of json sample.ts
""", "A small evidence bundle")
    add_body(doc, "Interpret common splits: if a local capture is clean but live playback is not, investigate network/reconnect/buffering. If VLC plays the live URL but Jellyfin transcodes or fails, inspect client/server negotiation. If both play but the guide is wrong, stop touching media settings and fix identity/time mapping. If recordings fail only during overlaps, inspect real padded reservations and source limits.")
    add_callout(doc, "Escalation package", "Share software versions, sanitized commands, exact UTC window, source type, ffprobe JSON, a short private sample when rights permit, packet-loss/interface counters, server playback reason, and the first relevant error - never the full secret URL.", fill=LIGHT_GOLD, accent=GOLD)

    chapter(doc, "A", "Quick-reference examples")
    doc.add_paragraph("A.1  Validation invariants", style="Heading 2")
    for item in [
        "Every enabled canonical channel has exactly one stable ID.",
        "Every exported M3U tvg-id maps to zero or one intended XMLTV channel@id; zero is flagged, more than one is an error.",
        "Every XMLTV programme references a declared channel and has start < stop when stop exists.",
        "No active programme overlaps another on the same channel unless the source explicitly models a clump/exception.",
        "Every stream route has a kind, priority, secret policy, timeout, health state, and last success/failure.",
        "Published M3U and XMLTV snapshots share a release/version identifier and pass cross-reference validation.",
        "Provider/tuner capacity is greater than the maximum planned padded overlap plus operational reserve.",
    ]:
        add_bullet(doc, item, bullet_id)
    doc.add_paragraph("A.2  FFmpeg decision shorthand", style="Heading 2")
    add_table(doc, ["Need", "First attempt", "Reason"], [
        ["Inspect", "ffprobe -show_programs -show_streams -of json", "Know what exists before selecting"],
        ["Record unchanged", "-map ... -c copy output.ts", "Preserve quality and reduce CPU"],
        ["Change container", "-map ... -c copy output.mkv/mp4/ts", "Remux only if target supports tracks"],
        ["Browser delivery", "HLS remux with copied tracks", "Low cost if keyframes/codecs are suitable"],
        ["Fix codec/bitrate/scale", "Decode/filter/encode chosen streams", "Only transcode the incompatible dimension"],
        ["Fix evidence", "Short capture + packet/timestamp probe", "Avoid speculative timestamp flags"],
    ], [2100, 4050, 3210], caption="Table A1. Start with the least transformative FFmpeg path.")
    doc.add_paragraph("A.3  Approximate live-latency intuition", style="Heading 2")
    add_body(doc, "These are engineering intuition, not protocol guarantees. Raw managed UDP/RTP can be sub-second to a few seconds; SRT/WebRTC can operate in sub-second-to-low-seconds ranges when configured and the network allows; low-latency HLS/DASH often lands in low single-digit seconds; conventional segmented HLS commonly reaches roughly 10-30 seconds. Encoding, GOP, segment publication, buffering, CDN, device, and recovery policy dominate the result.")

    chapter(doc, "B", "Glossary")
    glossary = [
        ("ABR", "Adaptive bitrate selection among multiple encoded representations."),
        ("Access module", "The player/server component that opens a file, URL, device, or protocol source."),
        ("Canonical channel ID", "A stable identifier owned by the integrating system, independent of display text and route."),
        ("CDM", "Content Decryption Module used by a DRM key system for protected playback."),
        ("Codec", "Algorithm and bitstream syntax that compresses one media track."),
        ("Container / mux", "Structure that packages one or more timed tracks and metadata."),
        ("DASH", "HTTP adaptive streaming described by an MPD and segmented representations."),
        ("Demux", "Read a container and separate its elementary streams/packets."),
        ("DTS", "Decode timestamp: when a compressed unit should be decoded."),
        ("DVB EIT", "DVB Event Information Table carrying present/following or schedule events."),
        ("EPG", "Electronic program guide: channels plus scheduled programme airings."),
        ("Elementary stream", "One encoded audio, video, subtitle, or data stream."),
        ("GOP", "Group of pictures; sequence around independently decodable keyframes."),
        ("HLS", "HTTP Live Streaming: master/media playlists and segments."),
        ("IGMP", "IPv4 multicast membership protocol between hosts and local multicast routers."),
        ("M3U / M3U8", "Line-oriented playlist; M3U8 conventionally denotes UTF-8 but may also be an HLS playlist."),
        ("MPEG-TS", "Packetized transport container designed for broadcast and streaming."),
        ("Multiplex / mux", "Several services/tracks interleaved in one transport; also the act of combining them."),
        ("PAT / PMT", "MPEG-TS tables that map programs to PMTs and PMTs to elementary stream PIDs."),
        ("PCR", "Program Clock Reference used to reconstruct the decoder's system time clock."),
        ("PID", "13-bit MPEG-TS packet identifier for a table or elementary stream."),
        ("PSIP", "ATSC tables describing virtual channels, time, ratings, and events."),
        ("PTS", "Presentation timestamp: when a decoded unit should be presented."),
        ("Remux", "Copy compressed packets into a different container or segment layout without re-encoding."),
        ("RTP / RTCP", "Real-time media transport plus reception and control reporting."),
        ("RTSP", "Session-control protocol commonly used to negotiate RTP delivery."),
        ("SRT", "Reliable low-latency transport over UDP with recovery and optional encryption."),
        ("Streamcopy", "FFmpeg packet copy selected with -c copy; no decode/filter/encode."),
        ("Timeshift", "Local rolling live buffer that enables pause and seek."),
        ("Transcode", "Decode and encode media into a different compressed representation."),
        ("tvg-id", "De facto M3U attribute commonly matched to XMLTV channel@id."),
        ("XMLTV", "XML interchange format for channels and programme listings."),
    ]
    for term, definition in glossary:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"{term}. ")
        r.bold = True
        p.add_run(definition)

    chapter(doc, "C", "Sources and research notes")
    add_body(doc, "Research method. Primary standards and official project documentation were preferred for technical claims. Context7 was used to locate current VLC, FFmpeg, and Jellyfin documentation/source; GitHub project documentation supplied real-world M3U/XMLTV and dynamic-channel conventions; Wikipedia was used only for broad orientation and cross-checking. Product behavior can change, so operational commands should be verified against the installed version.")
    add_body(doc, "Access date: 19 August 2026. Source labels [S#] used throughout the guide correspond to the entries below.", style="Small Text")
    for source_id, (title, url, note) in SOURCES.items():
        p = doc.add_paragraph(style="Source Text")
        r = p.add_run(f"[{source_id}] ")
        set_run_font(r, size=9, color=NAVY, bold=True)
        add_hyperlink(p, title, url)
        p.add_run(f" — {note}")

    # Mark the first row of every table for accessibility. For comparison
    # tables this is semantic; for one-row callout/cover furniture it keeps
    # assistive technology from treating the cells as an unlabeled data grid.
    for table in doc.tables:
        if table.rows:
            set_repeat_table_header(table.rows[0])

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(180)
    r = p.add_run("The shortest diagnosis is often the right boundary.")
    set_run_font(r, size=17, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Guide problem? Fix identity and time. Playback problem? Follow bytes, tracks, clocks, and capability.")
    set_run_font(r, size=11, color=MUTED, italic=True)

    return doc


def audit(doc: Document) -> None:
    assert len(doc.sections) == 1
    sec = doc.sections[0]
    assert abs(sec.page_width.inches - PAGE_W_IN) < 0.01
    assert abs(sec.left_margin.inches - MARGIN_IN) < 0.01
    assert doc.styles["Normal"].font.name == "Calibri"
    assert abs(doc.styles["Normal"].font.size.pt - 11) < 0.01
    assert len(doc.tables) >= 20
    assert len(doc.inline_shapes) == 4
    for table in doc.tables:
        # Footer table has 2 cols, body tables vary; all were explicitly sized.
        assert table._tbl.tblPr.find(qn("w:tblW")) is not None
        assert table._tbl.tblGrid is not None
    for para in doc.paragraphs:
        txt = para.text.lstrip()
        assert not txt.startswith("•")
        assert not (len(txt) > 2 and txt[0].isdigit() and txt[1:3] == ". ")
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "REDACTED" in body_text
    assert "PLACEHOLDER" not in body_text
    assert "turn0" not in body_text


if __name__ == "__main__":
    document = build_document()
    audit(document)
    document.save(OUT)
    print(f"Created {OUT}")
    print(f"Paragraphs: {len(document.paragraphs)}")
    print(f"Tables: {len(document.tables)}")
    print(f"Figures: {len(document.inline_shapes)}")
