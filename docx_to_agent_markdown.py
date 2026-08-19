from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.document import Document as _Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent
INPUT = ROOT / "IPTV_End_to_End_Field_Guide.docx"
OUTPUT = ROOT / "IPTV_End_to_End_Field_Guide.md"

IMAGE_TARGETS = {
    "image1.png": "iptv_guide_assets/architecture.png",
    "image2.png": "iptv_guide_assets/identity.png",
    "image3.png": "iptv_guide_assets/buffer_chain.png",
    "image4.png": "iptv_guide_assets/playback_decision.png",
}

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def iter_blocks(parent: _Document | _Cell):
    parent_element = parent.element.body if isinstance(parent, _Document) else parent._tc
    for child in parent_element.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def xml_bool(element, tag: str) -> bool:
    node = element.find(qn(tag))
    if node is None:
        return False
    value = node.get(qn("w:val"))
    return value not in {"0", "false", "off", "none"}


def run_text(run_element) -> str:
    parts: list[str] = []
    for node in run_element.iter():
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag == qn("w:tab"):
            parts.append("    ")
        elif node.tag in {qn("w:br"), qn("w:cr")}:
            parts.append("<br>")
    return "".join(parts)


def image_markdown(paragraph: Paragraph, run_element) -> str:
    blips = run_element.findall(f".//{{{A}}}blip")
    if not blips:
        return ""
    relationship_id = blips[0].get(f"{{{R}}}embed")
    if not relationship_id or relationship_id not in paragraph.part.rels:
        return ""
    target = Path(paragraph.part.rels[relationship_id].target_ref).name
    markdown_target = IMAGE_TARGETS.get(target, f"iptv_guide_assets/{target}")
    doc_props = run_element.findall(f".//{{{WP}}}docPr")
    alt = "IPTV guide diagram"
    if doc_props:
        alt = doc_props[0].get("descr") or doc_props[0].get("title") or alt
    alt = alt.replace("[", "(").replace("]", ")")
    return f"![{alt}]({markdown_target})"


def render_run(paragraph: Paragraph, run_element) -> str:
    image = image_markdown(paragraph, run_element)
    text = run_text(run_element)
    if not text:
        return image
    run_properties = run_element.find(qn("w:rPr"))
    if run_properties is not None:
        bold = xml_bool(run_properties, "w:b")
        italic = xml_bool(run_properties, "w:i")
        leading = text[: len(text) - len(text.lstrip())]
        trailing = text[len(text.rstrip()) :]
        core = text.strip()
        if core and bold and italic:
            text = f"{leading}***{core}***{trailing}"
        elif core and bold:
            text = f"{leading}**{core}**{trailing}"
        elif core and italic:
            text = f"{leading}*{core}*{trailing}"
    return f"{text}{image}"


def render_inline(paragraph: Paragraph) -> str:
    rendered: list[str] = []
    for child in paragraph._p:
        if child.tag == qn("w:r"):
            rendered.append(render_run(paragraph, child))
        elif child.tag == qn("w:hyperlink"):
            label = "".join(render_run(paragraph, run) for run in child.findall(qn("w:r")))
            relationship_id = child.get(qn("r:id"))
            if relationship_id and relationship_id in paragraph.part.rels:
                target = paragraph.part.rels[relationship_id].target_ref
                rendered.append(f"[{label}]({target})")
            else:
                rendered.append(label)
        elif child.tag in {qn("w:smartTag"), qn("w:sdt")}:
            for run in child.findall(f".//{{{W}}}r"):
                rendered.append(render_run(paragraph, run))
    return "".join(rendered).strip()


def numbering_format(document: _Document, paragraph: Paragraph) -> tuple[str, int] | None:
    paragraph_properties = paragraph._p.pPr
    if paragraph_properties is None or paragraph_properties.numPr is None:
        return None
    num_properties = paragraph_properties.numPr
    if num_properties.numId is None:
        return None
    num_id = str(num_properties.numId.val)
    level = int(num_properties.ilvl.val) if num_properties.ilvl is not None else 0
    root = document.part.numbering_part.element
    abstract_id = None
    for num in root.findall(qn("w:num")):
        if num.get(qn("w:numId")) == num_id:
            abstract = num.find(qn("w:abstractNumId"))
            abstract_id = abstract.get(qn("w:val")) if abstract is not None else None
            break
    if abstract_id is None:
        return "decimal", level
    for abstract in root.findall(qn("w:abstractNum")):
        if abstract.get(qn("w:abstractNumId")) != abstract_id:
            continue
        for level_node in abstract.findall(qn("w:lvl")):
            if int(level_node.get(qn("w:ilvl"), "0")) == level:
                num_format = level_node.find(qn("w:numFmt"))
                value = num_format.get(qn("w:val")) if num_format is not None else "decimal"
                return value, level
    return "decimal", level


def detect_code_language(code: str) -> str:
    stripped = code.lstrip()
    if stripped.startswith("<?xml"):
        return "xml"
    if stripped.startswith("channel_id:"):
        return "yaml"
    if stripped.startswith("#EXTM3U"):
        return "m3u"
    if stripped.startswith("{") or stripped.startswith("["):
        return "json"
    return "bash"


def normalize_heading(text: str) -> str:
    text = re.sub(r"^(\d+)\s{2,}", r"\1. ", text.strip())
    return re.sub(r"^([A-Z])\s{2,}", r"\1. ", text)


def cell_text(document: _Document, cell: _Cell) -> str:
    chunks: list[str] = []
    for paragraph in cell.paragraphs:
        text = render_inline(paragraph)
        if not text:
            continue
        numbering = numbering_format(document, paragraph)
        if numbering:
            fmt, _ = numbering
            marker = "-" if fmt == "bullet" else "1."
            text = f"{marker} {text}"
        chunks.append(text)
    return "<br>".join(chunks).replace("|", r"\|")


def render_table(document: _Document, table: Table) -> str:
    if len(table.rows) == 1 and len(table.columns) == 1:
        cell_paragraphs = table.cell(0, 0).paragraphs
        paragraphs = [render_inline(p) for p in cell_paragraphs]
        paragraphs = [p for p in paragraphs if p]
        if not paragraphs:
            return ""
        label = cell_paragraphs[0].text.strip() or paragraphs[0].strip("*")
        lines = [f"> **{label}**"]
        for paragraph in paragraphs[1:]:
            lines.extend([">", f"> {paragraph}"])
        return "\n".join(lines)

    rows = [[cell_text(document, cell) for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    header = rows[0]
    lines = ["| " + " | ".join(header) + " |"]
    lines.append("| " + " | ".join("---" for _ in header) + " |")
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def render_paragraph(document: _Document, paragraph: Paragraph) -> str:
    text = render_inline(paragraph)
    raw_text = paragraph.text.strip()
    if not text or text == "<br>":
        return ""
    style = paragraph.style.name

    if raw_text == "TECHNICAL FIELD GUIDE":
        return ""
    if raw_text == "IPTV, End to End":
        return "# IPTV, End to End"
    if raw_text == "How lineups, guide data, tuners, schedules, codecs, players, FFmpeg, VLC, and Jellyfin fit together":
        return f"> {raw_text}"
    if raw_text == "A practical reference for builders, operators, and technically curious readers":
        return f"*{raw_text}*"
    if raw_text == "Research edition • 19 August 2026":
        return "*Research edition — 19 August 2026*"

    if style == "Heading 1":
        return f"## {normalize_heading(raw_text)}"
    if style == "Heading 2":
        return f"### {normalize_heading(raw_text)}"
    if style == "Callout Label":
        return f"### {raw_text}"
    if style == "Code Block":
        code = paragraph.text.strip("\n")
        return f"```{detect_code_language(code)}\n{code}\n```"
    if style == "Caption":
        return f"*{text}*"
    if style == "Source Text":
        return f"- {text}"
    if style == "Small Text":
        return f"*{text}*"
    if style == "Lead":
        return f"> {text}"

    numbering = numbering_format(document, paragraph)
    if numbering:
        fmt, level = numbering
        marker = "-" if fmt == "bullet" else "1."
        return f"{'  ' * level}{marker} {text}"
    return text


def build_markdown(document: _Document) -> str:
    front_matter = """---
document_id: iptv-end-to-end-field-guide
title: "IPTV, End to End"
subtitle: "How lineups, guide data, tuners, schedules, codecs, players, FFmpeg, VLC, and Jellyfin fit together"
document_type: technical-field-guide
version: "1.0"
last_updated: 2026-08-19
language: en
audience:
  - builders
  - operators
  - technical readers
source_reference_format: "[S#]"
source_count: 42
topics:
  - IPTV architecture
  - XMLTV and EPG
  - network tuners
  - channel identity
  - extended M3U
  - scheduling and DVR
  - protocols and multicast
  - containers and codecs
  - VLC playback and buffering
  - FFmpeg inspection, remuxing, and transcoding
  - Jellyfin live TV and playback decisions
  - dynamic groups and virtual channels
  - observability, security, and troubleshooting
---"""

    parsing_notes = """## Parsing conventions

- Heading numbers are stable chapter identifiers.
- Claims cite bibliography entries using source IDs such as `[S12]`.
- Command examples use fenced code blocks with language identifiers.
- Tables use their first row as the header row.
- Diagram links use relative paths and include complete descriptive alt text.
- Example credentials, tokens, hostnames, and addresses are placeholders only."""

    blocks = [front_matter]
    inserted_notes = False
    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            rendered = render_paragraph(document, block)
            if rendered == "## Guide map" and not inserted_notes:
                blocks.extend([parsing_notes, rendered])
                inserted_notes = True
            elif rendered:
                blocks.append(rendered)
        else:
            rendered = render_table(document, block)
            if rendered:
                blocks.append(rendered)
    return "\n\n".join(blocks).rstrip() + "\n"


def main() -> None:
    document = Document(INPUT)
    OUTPUT.write_text(build_markdown(document), encoding="utf-8")
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
